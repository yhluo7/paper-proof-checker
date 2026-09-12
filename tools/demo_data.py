"""公开合成Demo的输入定义与生成逻辑。

同一份生成逻辑同时供 tools/generate_demo.py（写盘到 examples/quick-demo/）
和 tests/（写入临时目录）复用，保证Demo与测试输入永远一致。

Demo固定覆盖14项场景：
数值：完全匹配、百分比与小数换算、四舍五入兼容、P值阈值兼容、
      数值冲突、数值无来源、年份外的非结果数值、图片人工检查；
引用：直接支撑、部分支撑、结论冲突、已有全文但未找到证据、参考文献全文缺失；
映射：reference-map.json 手工映射与页码回溯。
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt

from tools.synth_pdf import build_text_pdf


# 固定文档时间戳，保证重复生成的DOCX字节稳定。
FIXED_TIMESTAMP = datetime(2026, 9, 12, 8, 0, 0, tzinfo=timezone.utc)

MANUSCRIPT_PARAGRAPHS: list[tuple[str, str]] = [
    ("heading", "Abstract"),
    ("body", "In this randomised trial, the intervention reduced mortality at one year."),
    ("heading", "Methods"),
    ("body", "All analyses were performed with version 4.2 of the statistics package."),
    ("heading", "Results"),
    ("body", "At the one-year follow-up, mortality in the intervention group was 32.6%."),
    ("body", "The hazard ratio for all-cause mortality was 0.72."),
    ("body", "The response rate in the intervention group was 0.458."),
    ("body", "The primary outcome met the prespecified significance criterion, P<0.05."),
    ("body", "Mortality in the usual care group was 41.2%."),
    ("body", "The absolute risk reduction was 6.3 percentage points."),
    ("body", "Adherence to the allocated treatment was 88.4% overall."),
    ("caption", "Table 1. Main outcomes"),
    ("table", "main"),
    ("heading", "Discussion"),
    ("body", "The survival estimates plotted in the supplementary figure show a long-term survival of 68.4%."),
    ("body", "Treatment reduced all-cause mortality at one year in the treated group [1]."),
    ("body", "The intervention improved survival in every prespecified subgroup [2]."),
    ("body", "The intervention reduced mortality compared with usual care [3]."),
    ("body", "The intervention reduced hospital admissions at one year [4]."),
    ("body", "The intervention reduced mortality in patients with severe disease [5]."),
    ("heading", "References"),
    ("body", "1. Smith J. Effects of the intervention on one-year mortality. Journal of Trials. 2022. doi:10.1234/demo.1"),
    ("body", "2. Jones K. Subgroup analysis of the intervention trial. Journal of Trials. 2023. doi:10.1234/demo.2"),
    ("body", "3. Brown L. Long-term outcomes after the intervention. Journal of Trials. 2021. doi:10.1234/demo.3"),
    ("body", "4. Chen M. Outcomes after the intervention: a one-year cohort study. Journal of Trials. 2020. doi:10.1234/demo.4"),
    ("body", "5. Miller R. Severe disease outcomes after the intervention. Journal of Trials. 2019. doi:10.1234/demo.5"),
]

MAIN_TABLE: list[list[str]] = [
    ["Outcome", "Rate", "HR", "P value"],
    ["Mortality at one year", "32.6%", "0.7234", "0.032"],
    ["Response rate", "45.8%", "0.85", "0.268"],
    ["Usual care", "38.9%", "1.05", "0.401"],
]

SUPPLEMENT_PARAGRAPHS: list[tuple[str, str]] = [
    ("heading", "Supplementary material"),
    ("caption", "Table S1. Supplementary outcome rates"),
]

SUPPLEMENT_TABLE: list[list[str]] = [
    ["Outcome", "Rate"],
    ["Adherence to allocated treatment", "88.4%"],
]

REFERENCE_PDFS: dict[str, dict] = {
    "Smith_2022.pdf": {
        "paragraphs": [
            "Effects of the intervention on one-year mortality",
            "doi:10.1234/demo.1",
            "Methods",
            "We randomly assigned adults admitted with severe disease to the intervention or to usual care.",
            "Results",
            "The intervention reduced all-cause mortality at one year in the treated group.",
            "Conclusion",
            "One-year survival is improved in the treated population.",
        ],
        "title": "Effects of the intervention on one-year mortality",
        "author": "Smith J",
    },
    "archive/Jones_2023.pdf": {
        "paragraphs": [
            "Subgroup analysis of the intervention trial",
            "doi:10.1234/demo.2",
            "Methods",
            "We repeated the main analysis within each prespecified subgroup.",
            "Results",
            "In the subgroup of patients with diabetes, the intervention improved survival at one year.",
            "No other prespecified subgroup showed a significant improvement.",
        ],
        "title": "Subgroup analysis of the intervention trial",
        "author": "Jones K",
    },
    "Brown_2021.pdf": {
        # 故意不写DOI，并让元数据标题与被引条目不同，
        # 用于验证 reference-map.json 的手工映射优先于自动匹配。
        "paragraphs": [
            "Internal analysis report",
            "Methods",
            "We compared one-year mortality between the intervention and usual care.",
            "Results",
            "The intervention did not reduce mortality compared with usual care at one year.",
        ],
        "title": "Internal analysis report",
        "author": "",
    },
    "archive/Chen_2020.pdf": {
        "paragraphs": [
            "Outcomes after the intervention: a one-year cohort study",
            "doi:10.1234/demo.4",
            "Methods",
            "Quality-of-life scores were the only outcome analysed in this cohort.",
            "Results",
            "The intervention did not change quality-of-life scores at one year.",
        ],
        "title": "Outcomes after the intervention: a one-year cohort study",
        "author": "Chen M",
    },
}

REFERENCE_MAP: dict[str, str] = {
    "1": "Smith_2022.pdf",
    "2": "archive/Jones_2023.pdf",
    "3": "Brown_2021.pdf",
    "4": "archive/Chen_2020.pdf",
}


def _configure_docx(document: Document) -> None:
    """固定文档时间戳并统一正文字号，保证重复生成结果稳定。"""

    document.core_properties.created = FIXED_TIMESTAMP
    document.core_properties.modified = FIXED_TIMESTAMP
    document.core_properties.author = "paper-proof-checker demo"
    document.core_properties.title = "paper-proof-checker synthetic demo"
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)


def _freeze_zip_timestamps(path: Path) -> None:
    """把DOCX内部的zip时间戳固定。

    python-docx用当前时间写入zip条目，导致同一份内容每次生成的字节都不同。
    公开Demo需要冻结基准，因此在保存后重写zip并统一时间戳。
    """

    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in entries:
            frozen = zipfile.ZipInfo(info.filename, date_time=(2026, 9, 12, 8, 0, 0))
            frozen.compress_type = zipfile.ZIP_DEFLATED
            frozen.external_attr = info.external_attr
            target.writestr(frozen, data)
    path.write_bytes(buffer.getvalue())


def build_manuscript(path: Path) -> Path:
    """生成主文DOCX。"""

    document = Document()
    _configure_docx(document)
    for kind, text in MANUSCRIPT_PARAGRAPHS:
        if kind == "heading":
            document.add_heading(text, level=1)
        elif kind == "caption":
            document.add_paragraph(text, style="Caption")
        elif kind == "table":
            table = document.add_table(rows=len(MAIN_TABLE), cols=len(MAIN_TABLE[0]))
            for row_index, row in enumerate(MAIN_TABLE):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = value
        else:
            document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    _freeze_zip_timestamps(path)
    return path


def build_supplement(path: Path) -> Path:
    """生成补充材料DOCX。"""

    document = Document()
    _configure_docx(document)
    for kind, text in SUPPLEMENT_PARAGRAPHS:
        if kind == "heading":
            document.add_heading(text, level=1)
        elif kind == "caption":
            document.add_paragraph(text, style="Caption")
        else:
            document.add_paragraph(text)
    table = document.add_table(rows=len(SUPPLEMENT_TABLE), cols=len(SUPPLEMENT_TABLE[0]))
    for row_index, row in enumerate(SUPPLEMENT_TABLE):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    _freeze_zip_timestamps(path)
    return path


def build_reference_pdfs(reference_dir: Path) -> list[Path]:
    """生成4篇带文本层的虚构参考文献PDF。"""

    written = []
    for relative, spec in REFERENCE_PDFS.items():
        written.append(
            build_text_pdf(
                reference_dir / relative,
                spec["paragraphs"],
                title=spec["title"],
                author=spec["author"],
            )
        )
    return written


def build_demo_project(root: Path) -> dict:
    """在root下生成完整Demo输入，返回关键路径。"""

    import json

    root = Path(root)
    reference_dir = root / "references"
    reference_dir.mkdir(parents=True, exist_ok=True)
    manuscript = build_manuscript(root / "manuscript.docx")
    supplement = build_supplement(root / "supplement.docx")
    pdfs = build_reference_pdfs(reference_dir)
    map_path = root / "reference-map.json"
    map_path.write_text(
        json.dumps(REFERENCE_MAP, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return {
        "root": root,
        "manuscript": manuscript,
        "supplement": supplement,
        "reference_dir": reference_dir,
        "reference_map": map_path,
        "pdfs": pdfs,
    }
