"""DOCX解析：引用展开、修订检测、表格单元格与题注识别。"""

from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from paper_proof_checker.docx_parser import extract_citations, has_tracked_changes, parse_docx
from tools.demo_data import build_manuscript, build_supplement


class DocxParserTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_citation_range_expansion(self) -> None:
        citations = extract_citations("Prior studies [1, 3–5] support this claim.")
        self.assertEqual(citations[0]["reference_ids"], [1, 3, 4, 5])

    def test_tracked_changes_are_detected(self) -> None:
        source = build_manuscript(self.root / "manuscript.docx")
        changed = self.root / "tracked.docx"
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(changed, "w") as target:
            for item in original.infolist():
                data = original.read(item.filename)
                if item.filename == "word/document.xml":
                    data = data.replace(b"<w:body>", b"<w:body><w:ins>", 1)
                target.writestr(item, data)
        self.assertTrue(has_tracked_changes(changed))

    def test_every_table_cell_becomes_exactly_one_block(self) -> None:
        """回归：曾用id(cell._tc)去重，导致部分单元格被误判为重复而丢失。"""

        path = build_manuscript(self.root / "manuscript.docx")
        parsed = parse_docx(path, "manuscript")
        cells = [block for block in parsed["blocks"] if block["kind"] == "table_cell"]
        # 主表为4行4列，共16个单元格。
        self.assertEqual(len(cells), 16)
        texts = [block["text"] for block in cells]
        for expected in ["32.6%", "0.7234", "0.032", "45.8%", "0.268", "38.9%", "0.401"]:
            self.assertIn(expected, texts)

    def test_table_cell_carries_column_header_and_row_context(self) -> None:
        path = build_manuscript(self.root / "manuscript.docx")
        parsed = parse_docx(path, "manuscript")
        cell = next(
            block
            for block in parsed["blocks"]
            if block["kind"] == "table_cell" and block["text"] == "0.032"
        )
        self.assertEqual(cell["column_header"], "P value")
        self.assertIn("Mortality at one year", cell["context_text"])
        self.assertEqual(cell["location"]["table_index"], 1)

    def test_table_caption_is_not_reported_as_figure_caption(self) -> None:
        path = build_manuscript(self.root / "manuscript.docx")
        parsed = parse_docx(path, "manuscript")
        kinds = {block["text"]: block["kind"] for block in parsed["blocks"]}
        self.assertEqual(kinds["Table 1. Main outcomes"], "table_caption")

    def test_numbered_references_and_citations_are_extracted(self) -> None:
        path = build_manuscript(self.root / "manuscript.docx")
        parsed = parse_docx(path, "manuscript")
        self.assertEqual([item["reference_id"] for item in parsed["references"]], [1, 2, 3, 4, 5])
        self.assertEqual(len(parsed["citations"]), 5)

    def test_numbers_mode_can_skip_vancouver_requirements(self) -> None:
        """没有编号参考文献与方括号引用时，只检查数值不应报错。"""

        path = self.root / "plain.docx"
        document = Document()
        document.add_heading("Results", level=1)
        document.add_paragraph("Mortality was 12.5%.")
        document.save(path)
        with self.assertRaises(ValueError):
            parse_docx(path, "manuscript")
        parsed = parse_docx(path, "manuscript", require_vancouver=False)
        self.assertEqual(parsed["references"], [])
        self.assertTrue(parsed["blocks"])

    def test_supplement_is_parsed_as_supplement_document(self) -> None:
        path = build_supplement(self.root / "supplement.docx")
        parsed = parse_docx(path, "supplement")
        self.assertTrue(all(block["document"] == "supplement" for block in parsed["blocks"]))
        source = next(
            block
            for block in parsed["blocks"]
            if block["kind"] == "table_cell" and block["text"] == "88.4%"
        )
        self.assertEqual(source["column_header"], "Rate")


if __name__ == "__main__":
    unittest.main()
